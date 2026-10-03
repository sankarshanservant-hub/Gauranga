"""Собирает перевод «Карчи» Говинды Даса в один Markdown, DOCX и PDF — для RU и EN.

Порядок: титул и предисловие переводчика, предисловие издателей 1926 г. (00a, 00b),
текст поэмы (01–12). Зависимости: pip install python-docx reportlab; шрифт Liberation Serif.
Запуск: python3 build.py [ru|en] (по умолчанию оба).
"""
import os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
FONT_DIR = '/usr/share/fonts/truetype/liberation'
DEVA_FONT = os.path.join(HERE, '..', 'fonts', 'NotoSerifDevanagari-Regular.ttf')

META = {
    'ru': dict(
        out='Govinda-Karcha-ru', intro=['00a-vvedenie.md', '00b-vvedenie.md'],
        title='«Карча» Говинды Даса', author='Говинда Дас (рус. пер.)',
        toc='Оглавление', notes='Примечания',
        preface="""# «Карча» Говинды Даса

*Дневник странствий Шри Чайтаньи Махапрабху по Южной и Западной Индии*

*Перевод с бенгальского на русский*

---

## От переводчика

«Карча» (кадача — «записи, дневник») приписывается Говинде Дасу, кузнецу из Канчананагара
под Бурдваном, который, по его словам, в 1509–1510 гг. ушёл из дома, стал слугой Шри Чайтаньи,
присутствовал при его санньясе в Катве и сопровождал его в двухлетнем странствии по югу и
западу Индии — до Каньякумари, Пандхарпура, Сомнатха и Двараки — и обратно в Пури.

Книга впервые издана в 1895 г. Джайгопалом Госвами из Шантипура. Её подлинность оспаривается:
многие исследователи и большинство гаудия-вайшнавских авторитетов считают её сочинением XIX в.;
её защитником был Динешчандра Сен, предисловие которого к изданию 1926 г. переведено здесь
полностью для ознакомления с аргументами одной из сторон.

**Источники перевода.** Основной текст — издание «Gobinda Daser Karcha» (224 с., без даты),
распознанное Tesseract и сверенное со вторым экземпляром того же издания и со сканами.
Пропуск на с. 116–117 и окончание восполнены по изданию Калькуттского университета 1926 г.
(под ред. Д. Ч. Сена и Банвари Лала Госвами), из которого переведено и предисловие.

**Принципы.** Двустишия-паяры переведены построчно, без рифмы, и пронумерованы сквозной
нумерацией (в оригинале нумерации нет); в квадратных скобках курсивом даны страницы основного
издания. Испорченные или утраченные строки оговорены в примечаниях переводчика в конце каждой
части. Имена и термины — по `GLOSSARY.md`.

---

## Оглавление

"""),
    'en': dict(
        out='Govinda-Karcha-en', intro=['00a-introduction.md', '00b-introduction.md'],
        title="Govinda Dasa's Karcha", author='Govinda Dasa (Eng. tr.)',
        toc='Contents', notes='Notes',
        preface="""# Govinda Dasa's Karcha

*A diary of Sri Chaitanya Mahaprabhu's travels in South and West India*

*Translated from the Bengali*

---

## Translator's Preface

The *Karcha* (*kadacha*, "notes, diary") is ascribed to Govinda Dasa, a smith of Kanchananagar
near Burdwan, who says that in 1509–10 he left home, became Sri Chaitanya's servant, witnessed
his sannyasa at Katwa and accompanied him on his two-year pilgrimage through the South and West
of India — to Kanyakumari, Pandharpur, Somnath and Dvaraka — and back to Puri.

The book was first published in 1895 by Jaygopal Goswami of Shantipur. Its authenticity is
disputed: many scholars and most Gaudiya Vaishnava authorities regard it as a nineteenth-century
composition. Its champion was Dinesh Chandra Sen, whose introduction to the 1926 edition is
translated here in full so that the reader may weigh one side's arguments.

**Sources.** The base text is the edition *Gobinda Daser Karcha* (224 pp., undated), OCR'd with
Tesseract and checked against a second copy of the same edition and against the scans. The lacuna
on pp. 116–117 and the ending are supplied from the Calcutta University edition of 1926 (ed.
D. C. Sen and Banwari Lal Goswami), from which the introduction is also translated.

**Method.** The payar couplets are translated line by line, without rhyme, and numbered
continuously (the original has no numbering); the page numbers of the base edition are given in
italic square brackets. Corrupt or missing lines are noted in the translator's notes at the end
of each part. Names and terms follow `GLOSSARY.md`.

---

## Contents

"""),
}


def files(lang):
    m = META[lang]
    return [os.path.join(HERE, lang, f) for f in m['intro']] + \
           [os.path.join(HERE, lang, f'{i:02d}.md') for i in range(1, 13)]


def build_md(lang):
    m = META[lang]
    toc, body = [], []
    for p in files(lang):
        key = os.path.basename(p).split('-')[0].split('.')[0]          # 00a, 00b, 01 …
        text = open(p, encoding='utf-8').read().strip()
        text = re.sub(r'<!--.*?-->', '', text, flags=re.S)             # служебные метки страниц
        text = re.sub(r'\[\^([^\]]+)\]', lambda mm: f'[^{key}-{mm.group(1)}]', text)  # уникальные id сносок
        title = text.splitlines()[0].lstrip('# ').strip()
        sub = re.search(r'^\*([^*\[].+)\*$', text, re.M)
        toc.append(f'- {title}' + (f' — {sub.group(1)}' if sub and key[:2] != '00' else ''))
        body.append(text)
    md = m['preface'] + '\n'.join(toc) + '\n\n---\n\n' + '\n\n---\n\n'.join(body) + '\n'
    open(os.path.join(HERE, m['out'] + '.md'), 'w', encoding='utf-8').write(md)
    return md


def parse(md):
    """Разбивает Markdown на блоки (вид, данные) по пустым строкам."""
    blocks = []
    for para in re.split(r'\n\s*\n', md):
        lines = [l.rstrip() for l in para.strip('\n').splitlines() if l.strip()]
        if not lines:
            continue
        first = lines[0]
        if m := re.match(r'^(#{1,4}) (.+)', first):
            blocks.append(('h%d' % min(len(m.group(1)), 3), m.group(2)))
            if len(lines) > 1:
                blocks.extend(parse('\n'.join(lines[1:])))
        elif first.strip() == '---':
            blocks.append(('hr', ''))
        elif all(l.lstrip().startswith('|') for l in lines):
            rows = [[c.strip() for c in l.strip().strip('|').split('|')] for l in lines
                    if not re.match(r'^\s*\|[\s:|-]+\|\s*$', l)]
            blocks.append(('table', rows))
        elif re.search(r'[\u0900-\u097F]', para) and not re.match(r'^\*\*\d', first) and not first.startswith('>'):
            blocks.append(('deva', lines))
        elif all(l.startswith('>>') for l in lines):
            blocks.append(('comment', [l.lstrip('>').strip() for l in lines]))
        elif all(l.startswith('>') for l in lines):
            blocks.append(('quote', [l.lstrip('>').strip() for l in lines]))
        elif re.match(r'^\[\^[^\]]+\]:', first):
            notes, cur = [], None
            for l in lines:
                if mm := re.match(r'^\[\^([^\]]+)\]:\s*(.*)', l):
                    cur = [mm.group(1).split('-', 1)[-1], mm.group(2)]
                    notes.append(cur)
                elif cur:
                    cur[1] += ' ' + l.strip()
            blocks.extend(('note', n) for n in notes)
        elif all(re.match(r'^\s*[-*] ', l) for l in lines):
            blocks.extend(('bullet', re.sub(r'^\s*[-*] ', '', l)) for l in lines)
        elif re.match(r'^\*\*\d+\.\*\*', first):
            blocks.append(('verse', lines))
        else:
            blocks.append(('p', lines))
    return blocks


def inline(text, html=True):
    """Markdown-разметка строки → HTML-подобная разметка reportlab."""
    s = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;') if html else text
    s = re.sub(r'\[\^[^\]-]+-([^\]]+)\]', r'<super>\1</super>', s)
    s = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', s)
    s = re.sub(r'(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])', r'<i>\1</i>', s)
    s = re.sub(r'`([^`]+)`', r'\1', s)
    return s


def docx_runs(par, text, size=None):
    from docx.shared import Pt
    for part in re.split(r'(<b>.*?</b>|<i>.*?</i>|<super>.*?</super>)', inline(text, html=False)):
        if not part:
            continue
        m = re.match(r'<(b|i|super)>(.*)</\1>', part)
        r = par.add_run(m.group(2) if m else part)
        if m:
            r.bold = m.group(1) == 'b'
            r.italic = m.group(1) == 'i'
            r.font.superscript = m.group(1) == 'super'
        if size:
            r.font.size = Pt(size)


def build_docx(blocks, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, Cm

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Liberation Serif'
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    first_h1 = True
    for kind, data in blocks:
        if kind == 'hr':
            continue
        if kind[0] == 'h':
            lvl = int(kind[1])
            if lvl == 1 and not first_h1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_h1 = first_h1 and lvl != 1
            h = doc.add_heading('', level=lvl)
            docx_runs(h, data)
            if lvl == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if kind == 'table':
            t = doc.add_table(rows=0, cols=max(len(r) for r in data))
            t.style = 'Table Grid'
            for i, row in enumerate(data):
                cells = t.add_row().cells
                for c, val in zip(cells, row):
                    docx_runs(c.paragraphs[0], f'**{val}**' if i == 0 and val else val, size=9)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        if kind == 'verse':
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            for i, l in enumerate(data):
                if i:
                    p.add_run().add_break()
                docx_runs(p, l)
        elif kind == 'quote':
            p.paragraph_format.left_indent = Cm(1.0)
            for i, l in enumerate(data):
                if i:
                    p.add_run().add_break()
                docx_runs(p, f'*{l}*' if l and not l.startswith('*') else l)
        elif kind == 'deva':
            from docx.oxml.ns import qn
            p.paragraph_format.left_indent = Cm(1.0)
            for i, l in enumerate(data):
                if i:
                    p.add_run().add_break()
                r = p.add_run(l)
                r.font.size = Pt(12)
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:cs'), 'Noto Serif Devanagari')
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Noto Serif Devanagari')
        elif kind == 'comment':
            p.paragraph_format.left_indent = Cm(1.0)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for i, l in enumerate(data):
                if i:
                    p.add_run().add_break()
                docx_runs(p, l, size=10)
        elif kind == 'note':
            p.paragraph_format.left_indent = Cm(0.5)
            docx_runs(p, f'**{data[0]}.** {data[1]}', size=9)
        elif kind == 'bullet':
            p.paragraph_format.left_indent = Cm(0.6)
            docx_runs(p, '• ' + data)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for i, l in enumerate(data):
                if i:
                    p.add_run().add_break()
                docx_runs(p, l)
    doc.save(path)


def build_pdf(blocks, path, title, author):
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (HRFlowable, PageBreak, Paragraph, SimpleDocTemplate,
                                    Table, TableStyle)
    from reportlab.lib.fonts import addMapping

    for name, f in [('Serif', 'LiberationSerif-Regular.ttf'), ('Serif-B', 'LiberationSerif-Bold.ttf'),
                    ('Serif-I', 'LiberationSerif-Italic.ttf'), ('Serif-BI', 'LiberationSerif-BoldItalic.ttf')]:
        pdfmetrics.registerFont(TTFont(name, os.path.join(FONT_DIR, f)))
    if os.path.exists(DEVA_FONT):
        pdfmetrics.registerFont(TTFont('Deva', DEVA_FONT))
    addMapping('Serif', 0, 0, 'Serif'); addMapping('Serif', 1, 0, 'Serif-B')
    addMapping('Serif', 0, 1, 'Serif-I'); addMapping('Serif', 1, 1, 'Serif-BI')

    base = ParagraphStyle('base', fontName='Serif', fontSize=11, leading=15)
    S = {
        'h1': ParagraphStyle('h1', base, fontSize=19, leading=24, alignment=TA_CENTER, spaceAfter=10),
        'h2': ParagraphStyle('h2', base, fontSize=13.5, leading=18, spaceBefore=8, spaceAfter=6),
        'h3': ParagraphStyle('h3', base, fontSize=12, leading=16, spaceBefore=6, spaceAfter=4),
        'p': ParagraphStyle('p', base, spaceAfter=5, alignment=TA_JUSTIFY),
        'verse': ParagraphStyle('verse', base, leftIndent=0.9 * cm, firstLineIndent=-0.9 * cm, spaceAfter=4),
        'quote': ParagraphStyle('quote', base, leftIndent=1.0 * cm, spaceAfter=5),
        'deva': ParagraphStyle('deva', base, fontName='Deva', fontSize=12.5, leading=21, leftIndent=1.0 * cm,
                               spaceAfter=3, shaping=1),
        'comment': ParagraphStyle('comment', base, fontSize=10, leading=13.5, leftIndent=1.0 * cm,
                                  spaceAfter=4, alignment=TA_JUSTIFY),
        'note': ParagraphStyle('note', base, fontSize=9, leading=12, leftIndent=0.5 * cm, spaceAfter=3,
                               alignment=TA_JUSTIFY),
        'bullet': ParagraphStyle('bullet', base, leftIndent=0.6 * cm, firstLineIndent=-0.35 * cm, spaceAfter=3),
        'cell': ParagraphStyle('cell', base, fontSize=8.5, leading=10.5),
    }

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Serif', 9)
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm, str(doc.page))
        canvas.restoreState()

    story, first_h1 = [], True
    for kind, data in blocks:
        if kind == 'hr':
            story.append(HRFlowable(width='100%', color='#999999', spaceBefore=6, spaceAfter=6))
        elif kind[0] == 'h':
            if kind == 'h1':
                if not first_h1:
                    story.append(PageBreak())
                first_h1 = False
            story.append(Paragraph(inline(data), S[kind]))
        elif kind == 'table':
            ncol = max(len(r) for r in data)
            rows = [[Paragraph(inline(c), S['cell']) for c in r] + [''] * (ncol - len(r)) for r in data]
            t = Table(rows, colWidths=[(A4[0] - 4.4 * cm) / ncol] * ncol, repeatRows=1)
            t.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
                                   ('BACKGROUND', (0, 0), (-1, 0), colors.whitesmoke),
                                   ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.append(t)
        elif kind == 'note':
            story.append(Paragraph(f'<b>{data[0]}.</b>&nbsp;' + inline(data[1]), S['note']))
        elif kind == 'bullet':
            story.append(Paragraph('•&nbsp;' + inline(data), S['bullet']))
        elif kind == 'quote':
            story.append(Paragraph('<br/>'.join(f'<i>{inline(l)}</i>' for l in data), S['quote']))
        else:
            story.append(Paragraph('<br/>'.join(inline(l) for l in data), S[kind]))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title=title, author=author)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == '__main__':
    for lang in (sys.argv[1:] or ['ru', 'en']):
        m = META[lang]
        blocks = parse(build_md(lang))
        build_docx(blocks, os.path.join(HERE, m['out'] + '.docx'))
        build_pdf(blocks, os.path.join(HERE, m['out'] + '.pdf'), m['title'], m['author'])
        print(lang, 'ok', len(blocks), 'blocks')
