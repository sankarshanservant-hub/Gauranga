"""Собирает главы 01.md–22.md в один Markdown, DOCX (python-docx) и PDF (reportlab).

Зависимости: pip install python-docx reportlab; шрифт Liberation Serif.
"""
import os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = 'Advaita-Prakash-ru'

PREFACE = """# Шри Адвайта-пракаша

*Ишана Нагара*

*Перевод с бенгальского на русский*

---

## От переводчика

«Адвайта-пракаша» — жизнеописание Шри Адвайты Ачарьи, одного из главных спутников
Шри Чайтаньи Махапрабху. Его составил Ишана Нагара — ученик и слуга Адвайты, выросший
в его доме. Как сказано в последней главе, книга закончена в 1490 году эры Шака
(1568 год н. э.) в Лауде.

Перевод сделан по изданию, где бенгальские двустишия-паяры набраны деванагари, а за
каждым следует современный хинди-пересказ. Здесь переведён **только оригинальный
текст Ишаны Нагары**, двустишие за двустишием, с сохранением нумерации. Хинди-пересказ
издателя служил лишь для сверки понимания. Вставки издателя не переводились: выдержки
из других сочинений и современные песни.

Текст издания распознан (OCR) с небольшими погрешностями. Где чтение было неясным,
это отмечено в примечаниях в конце главы. Санскритские и бенгальские имена переданы
в упрощённой русской транскрипции.

---

## Оглавление

"""

def chapters():
    return [os.path.join(HERE, f'{i:02d}.md') for i in range(1, 23)]

def build_md():
    toc, body = [], []
    for i, p in enumerate(chapters(), 1):
        text = open(p, encoding='utf-8').read().strip()
        title = text.splitlines()[0].lstrip('# ').strip()
        sub = re.search(r'^\*(.+)\*$', text, re.M)
        toc.append(f'{i}. {title}' + (f' — {sub.group(1)}' if sub else ''))
        body.append(text)
    md = PREFACE + '\n'.join(toc) + '\n\n---\n\n' + '\n\n---\n\n'.join(body) + '\n'
    open(os.path.join(HERE, OUT + '.md'), 'w', encoding='utf-8').write(md)
    return md

FONT_DIR = '/usr/share/fonts/truetype/liberation'


def parse(md):
    """Разбирает собранный Markdown на блоки: (вид, текст, номер)."""
    blocks = []
    for line in md.splitlines():
        if m := re.match(r'^(#{1,3}) (.+)', line):
            blocks.append(('h%d' % len(m.group(1)), m.group(2), None))
        elif m := re.match(r'^(\d+)\. (.+)', line):
            blocks.append(('verse', m.group(2), m.group(1)))
        elif m := re.match(r'^- (.+)', line):
            blocks.append(('bullet', m.group(1), None))
        elif line.strip() == '---':
            blocks.append(('hr', '', None))
        elif line.strip():
            blocks.append(('p', line, None))
    return blocks


def runs(text):
    """Делит строку на фрагменты (текст, жирный, курсив) по разметке ** и *."""
    out = []
    for part in re.split(r'(\*\*.+?\*\*|\*.+?\*)', text):
        if not part:
            continue
        if part.startswith('**'):
            out.append((part[2:-2], True, False))
        elif part.startswith('*') and len(part) > 1:
            out.append((part[1:-1], False, True))
        else:
            out.append((part.replace('`', ''), False, False))
    return out


def build_docx(blocks, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, Cm

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Liberation Serif'
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    first_h1 = True
    for kind, text, num in blocks:
        if kind == 'hr':
            continue
        if kind.startswith('h'):
            lvl = int(kind[1])
            if lvl == 1 and not first_h1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if lvl == 1:
                first_h1 = False
            h = doc.add_heading(text, level=lvl)
            if lvl == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        p = doc.add_paragraph()
        if kind == 'verse':
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(f'{num}. ')
            r.bold = True
        elif kind == 'bullet':
            p.paragraph_format.left_indent = Cm(0.6)
            p.add_run('• ')
        p.paragraph_format.space_after = Pt(3)
        for t, b, i in runs(text):
            r = p.add_run(t)
            r.bold, r.italic = b, i
    doc.save(path)


def build_pdf(blocks, path):
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (HRFlowable, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer)
    from reportlab.lib.fonts import addMapping

    for name, f in [('Serif', 'LiberationSerif-Regular.ttf'), ('Serif-B', 'LiberationSerif-Bold.ttf'),
                    ('Serif-I', 'LiberationSerif-Italic.ttf'), ('Serif-BI', 'LiberationSerif-BoldItalic.ttf')]:
        pdfmetrics.registerFont(TTFont(name, os.path.join(FONT_DIR, f)))
    addMapping('Serif', 0, 0, 'Serif'); addMapping('Serif', 1, 0, 'Serif-B')
    addMapping('Serif', 0, 1, 'Serif-I'); addMapping('Serif', 1, 1, 'Serif-BI')

    base = ParagraphStyle('base', fontName='Serif', fontSize=11, leading=15)
    styles = {
        'h1': ParagraphStyle('h1', base, fontSize=19, leading=24, alignment=TA_CENTER, spaceAfter=10),
        'h2': ParagraphStyle('h2', base, fontSize=13.5, leading=18, spaceBefore=8, spaceAfter=6),
        'h3': ParagraphStyle('h3', base, fontSize=12, leading=16, spaceBefore=6, spaceAfter=4),
        'p': ParagraphStyle('p', base, spaceAfter=5, alignment=TA_JUSTIFY),
        'verse': ParagraphStyle('verse', base, leftIndent=0.9 * cm, firstLineIndent=-0.9 * cm,
                                spaceAfter=3.5, alignment=TA_JUSTIFY),
        'bullet': ParagraphStyle('bullet', base, fontSize=9.5, leading=13, leftIndent=0.6 * cm,
                                 firstLineIndent=-0.35 * cm, spaceAfter=3),
    }

    def markup(text):
        s = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        s = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', s)
        s = re.sub(r'\*(.+?)\*', r'<i>\1</i>', s)
        return s.replace('`', '')

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Serif', 9)
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm, str(doc.page))
        canvas.restoreState()

    story, first_h1 = [], True
    for kind, text, num in blocks:
        if kind == 'hr':
            story.append(HRFlowable(width='100%', color='#999999', spaceBefore=6, spaceAfter=6))
            continue
        if kind == 'h1':
            if not first_h1:
                story.append(PageBreak())
            first_h1 = False
        if kind == 'verse':
            story.append(Paragraph(f'<b>{num}.</b>&nbsp;' + markup(text), styles['verse']))
        elif kind == 'bullet':
            story.append(Paragraph('•&nbsp;' + markup(text), styles['bullet']))
        else:
            story.append(Paragraph(markup(text), styles[kind]))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm,
                            title='Шри Адвайта-пракаша', author='Ишана Нагара (рус. пер.)')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == '__main__':
    blocks = parse(build_md())
    build_docx(blocks, os.path.join(HERE, OUT + '.docx'))
    build_pdf(blocks, os.path.join(HERE, OUT + '.pdf'))
    print('ok')
