#!/usr/bin/env python3
# Книга в Word (.docx): перевод шлок + «Амрита-праваха-бхашья» и «Анубхашья» (без бенгальского и пословного).
# Каждая глава — с новой страницы, название главы — «Заголовок 1» (навигация/оглавление в Word); сноски главы — в её конце.
# Использование: python3 tools/book_docx.py Adi-lila ru [full] Adi01 … Adi17   -> ed/docx/CC_Adi-lila-comm|full[-en].docx
#   full — полное издание: бенгальская шлока, транслитерация, пословный, перевод, оба комментария
import os, re, sys
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from md2pdf import verse_md, load_comm, SUB, has_bn, is_tl, TXT

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
name, lang, chaps = sys.argv[1], sys.argv[2], sys.argv[3:]
ED = 'full' if chaps and chaps[0] == 'full' else 'comm'
if chaps and chaps[0] in ('full', 'comm'): chaps = chaps[1:]
T = TXT[lang]
FONT = 'Times New Roman'
FONT_BN = 'Nirmala UI'          # есть в Windows; в других системах Word сам подставит бенгальский шрифт

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(148), Mm(210)                 # A5, как в PDF
sec.left_margin = sec.right_margin = Mm(14); sec.top_margin = Mm(14); sec.bottom_margin = Mm(16)
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.1
for hn, sz in (('Heading 1', 15), ('Heading 2', 11)):
    h = doc.styles[hn]; h.font.name = FONT; h.font.size = Pt(sz); h.font.color.rgb = RGBColor(0, 0, 0)

def runs(p, text, notes, chap, size=None, italic=False):
    """**жирный**, [^id] -> верхний индекс с номером сноски главы"""
    for part in re.split(r'(\*\*.+?\*\*|\[\^[\w-]+\])', text):
        if not part: continue
        m = re.match(r'\[\^([\w-]+)\]', part)
        if m:
            r = p.add_run(str(notes.index(m.group(1)) + 1) if m.group(1) in notes else '?'); r.font.superscript = True
        elif part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part)
        r.font.name = FONT; r.italic = italic or r.italic
        if size: r.font.size = Pt(size)

def bn_par(lines):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    for i, l in enumerate(lines):
        r = p.add_run(l.strip())
        rf = r._element.get_or_add_rPr().get_or_add_rFonts()
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), FONT_BN)
        r.font.size = Pt(12)
        sc = r._element.rPr.makeelement(qn('w:szCs'), {qn('w:val'): '24'}); r._element.rPr.append(sc)
        if i < len(lines) - 1: r.add_break()

def tlww_par(lines):
    ls = [x.strip() for x in lines if x.strip()]
    nodia = lang == 'ru' and not any(is_tl(x) for x in ls)   # транслитерация без диакритики (Мадхья 6.122)
    for k, l in enumerate(ls):
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
        p.paragraph_format.left_indent = Mm(4)
        tl = ('  ' not in l) if nodia else is_tl(l) if lang == 'ru' else k % 2 == 0
        p.paragraph_format.space_after = Pt(0 if tl else 2)
        r = p.add_run(l if tl else l.replace('  ', '\u2002\u2002'))
        r.font.name = FONT
        if tl: r.italic = True; r.font.size = Pt(10)
        else: r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def comm(text, notes, chap, label):
    first = True
    for para in re.split(r'\n\s*\n', text):
        para = para.strip()
        if not para or SUB.match(para): continue
        lines = [l.strip() for l in para.split('\n') if l.strip()]
        items = [l[2:] for l in lines] if all(l.startswith('- ') for l in lines) else [' '.join(lines)]
        for it in items:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if first:
                r = p.add_run(label + '. '); r.bold = True; r.italic = True; r.font.size = Pt(10); first = False
            runs(p, it, notes, chap, size=10)

for ci, chap in enumerate(chaps):
    md = verse_md(chap, lang)
    body, _, foot = md.partition('\n---\n')
    apb, fa = load_comm(chap, 'apb', lang); anu, fn = load_comm(chap, 'anu', lang)
    foot = foot + '\n' + fa + '\n' + fn
    notes = re.findall(r'^\[\^([\w-]+)\]:', foot, re.M)
    paras = [p for p in body.split('\n\n') if p.strip()]
    hdr = [l.strip().strip('*').strip() for l in paras[0].split('\n')]
    if ci:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for k, txt in enumerate(hdr[:5]):
        if k == 3:
            p = doc.add_heading(level=1); runs(p, '%s. %s' % (hdr[2], txt), notes, chap)
        elif k == 2: continue
        else:
            p = doc.add_paragraph(); runs(p, txt, notes, chap, size=10 if k == 4 else 11, italic=(k == 4))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if apb.get('0'):
        p = doc.add_heading(T['apb'], level=2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        comm(apb['0'], notes, chap, '')
    pend = []                                     # бенгальский и транслит+пословный ждут перевода: подзаголовок — над ними
    for p0 in paras[1:]:
        if has_bn(p0):
            if ED == 'full': pend.append(('bn', p0.split('\n')))
            continue
        m = re.match(r'^\((\d+(?:–\d+)?)\) (.*)', p0, re.S)
        if not m:
            if ED == 'full': pend.append(('tl', p0.split('\n')))
            continue
        lab = m.group(1); a, b = (lab.split('–') + [lab])[:2]
        keys = list(dict.fromkeys(k for k in list(apb) + list(anu) if k != '0' and int(a) <= int(k.split('–')[-1]) <= int(b)))
        sh = ''
        for k in keys:
            mm = SUB.search(anu.get(k, ''))
            if mm: sh = mm.group(1).strip()
        if sh:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
            runs(p, sh, notes, chap, size=10.5, italic=True)
        for i, (kind, ls) in enumerate(pend):
            if kind == 'bn':
                bn_par(ls)
                if i == 0 and not sh: doc.paragraphs[-1].paragraph_format.space_before = Pt(8)
            else: tlww_par(ls)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not sh and not pend: p.paragraph_format.space_before = Pt(8)
        elif pend: p.paragraph_format.space_before = Pt(3)
        pend = []
        r = p.add_run(lab + ' '); r.bold = True
        runs(p, ' '.join(m.group(2).split('\n')), notes, chap)
        for label, src in ((T['apb'], apb), (T['anu'], anu)):
            for k in keys:
                if k in src: comm(src[k], notes, chap, label)
    if notes:
        p = doc.add_heading(T['notes'], level=2)
        for i, k in enumerate(notes, 1):
            t = re.search(r'^\[\^%s\]:\s*(.*)$' % re.escape(k), foot, re.M).group(1)
            p = doc.add_paragraph(); r = p.add_run('%d. ' % i); r.font.size = Pt(9.5)
            runs(p, t, notes, chap, size=9.5)

od = os.path.join(ROOT, 'ed', 'docx'); os.makedirs(od, exist_ok=True)
out = os.path.join(od, 'CC_%s-%s%s.docx' % (name, ED, '' if lang == 'ru' else '-en'))
doc.save(out); print(out)
