import sys,re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
def setfont(run,beng):
    run.font.name='Noto Serif Bengali' if beng else 'Times New Roman'
    rpr=run._element.get_or_add_rPr(); rf=rpr.get_or_add_rFonts()
    for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a),run.font.name)
src,out=sys.argv[1],sys.argv[2]
doc=Document(); st=doc.styles['Normal']; st.font.size=Pt(11)
for block in re.split(r'\n\s*\n',open(src,encoding='utf-8').read()):
    lines=[l.rstrip() for l in block.strip('\n').split('\n')]
    if not any(lines): continue
    p=doc.add_paragraph()
    for i,l in enumerate(lines):
        if i: p.add_run().add_break()
        beng=bool(re.search('[ঀ-৿]',l))
        for j,part in enumerate(re.split(r'\*\*',l)):
            if not part: continue
            r=p.add_run(part); r.bold=(j%2==1); setfont(r,beng)
            if beng: r.font.size=Pt(12)
doc.save(out)
